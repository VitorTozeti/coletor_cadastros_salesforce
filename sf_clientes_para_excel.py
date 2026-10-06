"""Salesforce (Account) -> PowerApp_Clientes.xlsx na pasta do SharePoint sincronizada pelo OneDrive.

Roda no servidor do Caetano, 1x/dia. Só ADICIONA CardCodes que ainda não estão na tabela; linhas
existentes nunca são alteradas. Faz backup antes de gravar.

Uso:  python sf_clientes_para_excel.py --xlsx "<caminho local do PowerApp_Clientes.xlsx>" [--dry-run]

.env ao lado do script (nunca versionar): SF_CLIENT_ID, SF_CLIENT_SECRET, SF_LOGIN_URL.
"""
import argparse
import os
import re
import shutil
import sys
import uuid
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import load_workbook
from openpyxl.worksheet.table import Table

AQUI = Path(__file__).parent
COLUNAS = ["CardCode", "CardName", "Cod_Nome", "MailZipCod", "E_Mail", "MailStrNo", "Telefone", "Tipo_de_Conta",
           "Password", "__PowerAppsId__"]
# Tipo_de_Conta = CPF|CNPJ; Password = o próprio CPF/CNPJ formatado (padrão da planilha); __PowerAppsId__ = uuid4.

SOQL_CONTAS = (
    "SELECT Id, Name, CA_CodigoSAP__c, CA_Email__c, CA_EmailEmpresa__c, CA_Telefone1__c, CA_Celular__c, Phone, CA_CPF__c, CA_CNPJ__c "
    "FROM Account WHERE CA_CodigoSAP__c != null"
)
# Endereço: mais recente ATIVO, preferindo entrega (S). CA_TipoEndereco__c guarda 'S'/'B', não os rótulos.
SOQL_END = (
    "SELECT CA_Conta__c, CA_CEP__c, CA_Numero__c, CA_TipoEndereco__c, CreatedDate "
    "FROM CA_Endereco__c WHERE CA_StatusEndereco__c = 'Ativo' AND CA_Conta__c != null"
)


def carrega_env(caminho=AQUI / ".env"):
    for linha in caminho.read_text(encoding="utf-8").splitlines():
        if "=" in linha and not linha.lstrip().startswith("#"):
            k, v = linha.split("=", 1)
            os.environ.setdefault(k.strip(), v.split("#")[0].strip())


def sf_query(soql):
    r = requests.post(
        os.environ["SF_LOGIN_URL"] + "/services/oauth2/token",
        data={"grant_type": "client_credentials", "client_id": os.environ["SF_CLIENT_ID"],
              "client_secret": os.environ["SF_CLIENT_SECRET"]}, timeout=60)
    r.raise_for_status()
    t = r.json()
    h = {"Authorization": "Bearer " + t["access_token"]}
    url, params = t["instance_url"] + "/services/data/v60.0/query", {"q": soql}
    while url:
        d = requests.get(url, params=params, headers=h, timeout=120).json()
        yield from d["records"]
        url, params = (t["instance_url"] + d["nextRecordsUrl"], None) if not d["done"] else (None, None)


def melhor_endereco():
    """conta_id -> (cep, numero): ativo mais recente, entrega (S) antes de cobrança."""
    mapa = {}
    for e in sf_query(SOQL_END):
        chave = (e["CA_TipoEndereco__c"] == "S", e["CreatedDate"])
        atual = mapa.get(e["CA_Conta__c"])
        if atual is None or chave > atual[0]:
            mapa[e["CA_Conta__c"]] = (chave, (e["CA_CEP__c"] or "", e["CA_Numero__c"] or ""))
    return {k: v[1] for k, v in mapa.items()}


def documento(a):
    """(tipo, documento formatado) a partir de CA_CPF__c / CA_CNPJ__c; (None, None) se não tiver."""
    cpf, cnpj = re.sub(r"\D", "", a["CA_CPF__c"] or ""), re.sub(r"\D", "", a["CA_CNPJ__c"] or "")
    if len(cpf) == 11:
        return "CPF", f"{cpf[:3]}.{cpf[3:6]}.{cpf[6:9]}-{cpf[9:]}"
    if len(cnpj) == 14:
        return "CNPJ", f"{cnpj[:2]}.{cnpj[2:5]}.{cnpj[5:8]}/{cnpj[8:12]}-{cnpj[12:]}"
    return None, None


def linhas_salesforce():
    ends = melhor_endereco()
    for a in sf_query(SOQL_CONTAS):
        cep, num = ends.get(a["Id"], ("", ""))
        tipo, doc = documento(a)
        if not doc:  # sem CPF/CNPJ não há Password -> o cliente não conseguiria entrar no app; fica de fora
            continue
        cod, nome = a["CA_CodigoSAP__c"].strip(), (a["Name"] or "").strip()
        yield {
            "CardCode": cod, "CardName": nome, "Cod_Nome": f"{cod} - {nome}",
            "MailZipCod": cep, "E_Mail": a["CA_Email__c"] or a["CA_EmailEmpresa__c"] or "",
            "MailStrNo": num, "Telefone": a["CA_Telefone1__c"] or a["CA_Celular__c"] or a["Phone"] or "",
            "Tipo_de_Conta": tipo, "Password": doc, "__PowerAppsId__": str(uuid.uuid4()),
        }


def anexa_no_excel(xlsx, novas_fn, dry_run):
    wb = load_workbook(xlsx)
    ws = next((s for s in wb.worksheets if s.tables), None)
    if ws is None:
        sys.exit("A planilha não tem nenhuma TABELA do Excel (Inserir > Tabela). Abortado sem gravar.")
    tabela = next(iter(ws.tables.values()))
    c1, r1, c2, r2 = _limites(tabela.ref)
    cab = [str(ws.cell(r1, c).value).strip() for c in range(c1, c2 + 1)]
    faltam = [c for c in COLUNAS if c not in cab]
    if faltam:
        sys.exit(f"Colunas ausentes na tabela {tabela.name}: {faltam}. Cabeçalho: {cab}. Abortado sem gravar.")
    col = {nome: c1 + i for i, nome in enumerate(cab)}
    existentes = {str(ws.cell(r, col["CardCode"]).value or "").strip() for r in range(r1 + 1, r2 + 1)}
    novas = [l for l in novas_fn() if l["CardCode"] not in existentes]
    print(f"tabela {tabela.name}: {r2 - r1} linhas existentes; {len(novas)} CardCodes novos")
    if dry_run or not novas:
        return len(novas)
    backup = Path(xlsx).with_name(f"{Path(xlsx).stem}.bak-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    shutil.copy2(xlsx, backup)
    for i, linha in enumerate(novas, start=1):
        for nome, valor in linha.items():
            ws.cell(r2 + i, col[nome]).value = valor
    tabela.ref = f"{ws.cell(r1, c1).coordinate}:{ws.cell(r2 + len(novas), c2).coordinate}"
    wb.save(xlsx)
    return len(novas)


def _limites(ref):
    from openpyxl.utils import range_boundaries
    return range_boundaries(ref)


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--dry-run", action="store_true")
    a = ap.parse_args()
    carrega_env()
    try:
        n = anexa_no_excel(a.xlsx, lambda: list(linhas_salesforce()), a.dry_run)
    except PermissionError:
        sys.exit("Excel em uso/bloqueado pelo OneDrive — tente de novo mais tarde. Nada foi gravado.")
    print("dry-run: nada gravado" if a.dry_run else f"{n} linhas adicionadas")
